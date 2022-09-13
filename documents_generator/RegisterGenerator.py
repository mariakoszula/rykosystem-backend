from datetime import datetime
from documents_generator.DocumentGenerator import DocumentGenerator
from helpers.config_parser import config_parser
from helpers.date_converter import DateConverter
from models.contracts import ContractModel, AnnexModel
from models.program import ProgramModel
from helpers.file_folder_creator import DirectoryCreator
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx import Document


class RegisterGenerator(DocumentGenerator):
    CELL_TO_MERGE_MARK = "MERGE"

    def prepare_data(self):
        self.__prepare_school_data()
        self._document.merge_rows('no', self.records_to_merge)
        self._document.merge(
            semester_no=self.program.get_current_semester(),
            school_year=self.program.school_year,
            date=self.date
        )

    def __init__(self, program: ProgramModel):
        self.date = datetime.today().strftime('%d-%m-%Y')
        self.program = program
        self.records_to_merge = []
        self.contracts = ContractModel.all_filtered_by_program(self.program.id)
        DocumentGenerator.__init__(self,
                                   config_parser.get('DocTemplates', 'register'),
                                   DirectoryCreator.get_main_dir(school_year=self.program.school_year,
                                                                 semester_no=self.program.semester_no),
                                   config_parser.get('DocNames', 'register').format(self.date))

    def __prepare_school_data(self):
        for contract in sorted(self.contracts, key=lambda c: int(c.contract_no)):
            self.records_to_merge.append(RegisterGenerator.__prepare_contract_data(contract))

            for annex in contract.annex:
                self.records_to_merge.append(RegisterGenerator.__prepare_annex_data(annex))

    @staticmethod
    def __prepare_contract_data(contract: ContractModel):
        record_dict = dict()
        record_dict['no'] = str(contract.contract_no)
        record_dict['contract_info'] = f"{contract.contract_no}/{contract.contract_year}"
        record_dict['school_name'] = contract.school.name
        record_dict['school_nip'] = contract.school.nip
        record_dict['school_address'] = contract.school.address
        record_dict['school_city'] = contract.school.city
        record_dict['school_regon'] = contract.school.regon
        record_dict['school_phone'] = contract.school.phone
        record_dict['school_email'] = contract.school.email
        record_dict['kids_milk'] = str(contract.dairy_products)
        record_dict['kids_fruitveg'] = str(contract.fruitVeg_products)
        return record_dict

    @staticmethod
    def __prepare_annex_data(annex: AnnexModel):
        record_dict = dict()
        record_dict['no'] = RegisterGenerator.CELL_TO_MERGE_MARK
        record_dict['school_name'] = RegisterGenerator.CELL_TO_MERGE_MARK
        record_dict['school_nip'] = RegisterGenerator.CELL_TO_MERGE_MARK
        record_dict['school_address'] = RegisterGenerator.CELL_TO_MERGE_MARK
        record_dict['school_city'] = RegisterGenerator.CELL_TO_MERGE_MARK
        record_dict['school_regon'] = RegisterGenerator.CELL_TO_MERGE_MARK
        record_dict['school_phone'] = RegisterGenerator.CELL_TO_MERGE_MARK
        record_dict['school_email'] = RegisterGenerator.CELL_TO_MERGE_MARK
        record_dict['contract_info'] = RegisterGenerator.CELL_TO_MERGE_MARK
        validity_date = DateConverter.convert_date_to_string(annex.validity_date)
        record_dict['change_info'] = f"Aneks_{annex.no} {validity_date}*"
        record_dict['kids_milk'] = str(annex.dairy_products)
        record_dict['kids_fruitveg'] = str(annex.fruitVeg_products)
        return record_dict

    def generate(self, gen_pdf=True) -> None:
        DocumentGenerator.generate(self, gen_pdf=False)
        generated_docx = self.generated_documents[0].name
        RegisterGenerator.__merge_cells(generated_docx)
        self._generate_pdf(generated_docx)

    @staticmethod
    def __merge_cells(file_with_table_to_merge):
        document = Document(file_with_table_to_merge)
        for table in document.tables:
            for col in range(0, len(table.columns)):
                for row in range(0, len(table.rows)):
                    if table.cell(row, col).text == RegisterGenerator.CELL_TO_MERGE_MARK:
                        table.cell(row, col).text = ""
                        merged = table.cell(row - 1, col).merge(table.cell(row, col))
                        merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        table.allow_autofit = True
        document.save(file_with_table_to_merge)
